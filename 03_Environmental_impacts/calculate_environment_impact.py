from __future__ import annotations

import argparse
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


IMPACT_START_COL = 7


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip().lower())


def normalize_ingredient(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def recipe_key(country: Any, recipe_name: Any) -> str:
    return f"{normalize_text(country)}||{normalize_text(recipe_name)}"


def to_float(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def replace_code_suffix(code: str, suffix: str) -> str:
    if not code or not suffix or len(code) < 3:
        return code
    return code[:-3] + suffix.strip().upper()


def header_map(row: tuple[Any, ...]) -> dict[str, int]:
    return {str(value).strip(): i for i, value in enumerate(row) if value not in (None, "")}


def vec(length: int) -> list[float]:
    return [0.0] * length


def append_rows(ws, rows: list[list[Any]]) -> None:
    for row in rows:
        ws.append(row)
    if rows:
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"


def read_lcia(path: Path) -> tuple[list[str], list[str], dict[str, list[float]], dict[str, str]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    row4 = [cell.value for cell in ws[4]]
    row5 = [cell.value for cell in ws[5]]
    impact_cols = list(range(IMPACT_START_COL, ws.max_column + 1))
    impact_names: list[str] = []
    impact_units: list[str] = []
    for col in impact_cols:
        name = row4[col - 1] or f"impact_{col}"
        unit = row5[col - 1] or ""
        impact_names.append(str(name).strip())
        impact_units.append(str(unit).strip())

    lcia_by_code: dict[str, list[float]] = {}
    product_by_code: dict[str, str] = {}
    for row in ws.iter_rows(min_row=6, values_only=True):
        code = row[0]
        if not code:
            continue
        code = str(code).strip()
        values = vec(len(impact_names))
        for i, col in enumerate(impact_cols):
            val = to_float(row[col - 1] if col - 1 < len(row) else None)
            if val is not None:
                values[i] = val
        lcia_by_code[code] = values
        product_by_code[code] = "" if row[1] is None else str(row[1])
    wb.close()
    return impact_names, impact_units, lcia_by_code, product_by_code


def read_match(path: Path) -> tuple[dict[str, str], dict[str, str], dict[str, str], int]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    headers = header_map(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
    name_idx = headers["name"]
    code_idx = 1
    product_idx = 2
    code_by_ingredient: dict[str, str] = {}
    product_by_ingredient: dict[str, str] = {}
    canonical_by_ingredient: dict[str, str] = {}
    duplicate_names = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = normalize_ingredient(row[name_idx])
        code = row[code_idx]
        if not name or not code:
            continue
        if name in code_by_ingredient:
            duplicate_names += 1
            continue
        code_by_ingredient[name] = str(code).strip()
        product_by_ingredient[name] = "" if row[product_idx] is None else str(row[product_idx])
        canonical_by_ingredient[name] = "" if row[name_idx] is None else str(row[name_idx])
    wb.close()
    return code_by_ingredient, product_by_ingredient, canonical_by_ingredient, duplicate_names


def resolve_idea_code(base_code: str, country_code: str, lcia_by_code: dict[str, list[float]]) -> tuple[str | None, str]:
    country_candidate = replace_code_suffix(base_code, country_code)
    glo_candidate = replace_code_suffix(base_code, "GLO")
    if country_candidate in lcia_by_code:
        return country_candidate, "country_specific"
    if glo_candidate in lcia_by_code:
        return glo_candidate, "GLO_fallback"
    if base_code in lcia_by_code:
        return base_code, "base_code_fallback"
    return None, "IDEA_code_not_in_LCIA"


INGREDIENT_ALIASES = {
    "all purpose flour": "all purpose flour",
    "chicken thigh leg": "chicken",
    "chilli pepper green": "chilli pepper",
    "chilli pepper red": "red pepper",
    "cinnamon powder": "cinnamon",
    "cooking oil olive": "olive oil",
    "green pepper": "green bell pepper",
    "ground black pepper": "black pepper",
    "paprika powder": "paprika",
    "tomato sauce paste": "tomato paste",
}


def calculate(args: argparse.Namespace) -> None:
    base_dir = Path(args.base_dir).resolve()
    nutrition_path = Path(args.nutrition_path) if args.nutrition_path else base_dir / "Asia recipe data 2026" / "Asia_recipe_nutrition.xlsx"
    ingredient_info_path = Path(args.ingredient_info_path) if args.ingredient_info_path else base_dir / "Asia recipe data 2026" / "Ingredient_Info.xlsx"
    match_path = Path(args.match_path) if args.match_path else base_dir / "Environment Data" / "match.xlsx"
    lcia_path = Path(args.lcia_path) if args.lcia_path else base_dir / "Environment Data" / "LCIA.xlsx"
    output_path = Path(args.output_path) if args.output_path else base_dir / "Asia recipe data 2026" / "Asia_recipe_environment_impact.xlsx"

    for path in (nutrition_path, ingredient_info_path, match_path, lcia_path):
        if not path.exists():
            raise FileNotFoundError(path)

    print("Reading LCIA intensities...")
    impact_names, impact_units, lcia_by_code, product_by_code = read_lcia(lcia_path)
    impact_count = len(impact_names)

    print("Reading ingredient-to-IDEA matches...")
    code_by_ingredient, _product_by_ingredient, canonical_by_ingredient, duplicate_match_names = read_match(match_path)

    print("Reading recipe ingredient weights...")
    dish_agg: dict[str, dict[str, Any]] = {}
    dish_agg_by_source_id: dict[str, str] = {}
    dish_agg_by_name: dict[str, str] = {}
    ambiguous_name_keys: set[str] = set()

    ingredient_detail_rows: list[list[Any]] = [[
        "source_row",
        "source_idrecipe",
        "region",
        "countrycode",
        "country",
        "recipe_name",
        "ingredient_name",
        "ingredient_match_name",
        "gram",
        "IDEA_code_base",
        "IDEA_code_used",
        "match_status",
        "IDEA_product_name",
        *[f"impact_total_{i + 1}_{impact_names[i]} [{impact_units[i]}]" for i in range(impact_count)],
    ]]
    unmatched_rows: list[list[Any]] = [[
        "source_row",
        "source_idrecipe",
        "region",
        "countrycode",
        "country",
        "recipe_name",
        "ingredient_name",
        "gram",
        "IDEA_code_base",
        "country_code_candidate",
        "GLO_candidate",
        "reason",
    ]]

    ingredient_record_count = 0
    matched_ingredient_count = 0
    unmatched_ingredient_count = 0
    invalid_gram_count = 0
    alias_resolved_count = 0

    wb_ing = load_workbook(ingredient_info_path, data_only=True, read_only=True)
    ws_ing = wb_ing.active
    ing_headers = header_map(next(ws_ing.iter_rows(min_row=1, max_row=1, values_only=True)))
    for excel_row, row in enumerate(ws_ing.iter_rows(min_row=2, values_only=True), start=2):
        source_id = row[ing_headers["idrecipe"]]
        region = row[ing_headers["region"]]
        country_code = "" if row[ing_headers["countrycode"]] is None else str(row[ing_headers["countrycode"]])
        country = row[ing_headers["country"]]
        recipe_name = row[ing_headers["recipename"]]
        ingredient_name = row[ing_headers["name"]]
        gram = to_float(row[ing_headers["gram"]])

        if not normalize_text(country) or not normalize_text(recipe_name):
            continue

        name_key = recipe_key(country, recipe_name)
        source_id_key = normalize_text(source_id)
        agg_key = f"id::{source_id_key}" if source_id_key else f"name::{name_key}"
        if agg_key not in dish_agg:
            dish_agg[agg_key] = {
                "total": vec(impact_count),
                "ingredient_rows": 0,
                "matched_rows": 0,
                "unmatched_rows": 0,
                "source_ids": set(),
            }

        if source_id_key:
            dish_agg_by_source_id.setdefault(source_id_key, agg_key)
        if name_key not in dish_agg_by_name:
            dish_agg_by_name[name_key] = agg_key
        elif dish_agg_by_name[name_key] != agg_key:
            ambiguous_name_keys.add(name_key)

        agg = dish_agg[agg_key]
        agg["source_ids"].add("" if source_id is None else str(source_id))
        agg["ingredient_rows"] += 1
        ingredient_record_count += 1

        if gram is None:
            invalid_gram_count += 1
            unmatched_ingredient_count += 1
            agg["unmatched_rows"] += 1
            unmatched_rows.append([excel_row, source_id, region, country_code, country, recipe_name, ingredient_name, None, "", "", "", "invalid_or_blank_gram"])
            continue

        ing_key = normalize_ingredient(ingredient_name)
        lookup_key = ing_key
        alias_used = False
        if lookup_key not in code_by_ingredient and lookup_key in INGREDIENT_ALIASES:
            alias_target = normalize_ingredient(INGREDIENT_ALIASES[lookup_key])
            if alias_target in code_by_ingredient:
                lookup_key = alias_target
                alias_used = True
                alias_resolved_count += 1

        if lookup_key not in code_by_ingredient:
            unmatched_ingredient_count += 1
            agg["unmatched_rows"] += 1
            unmatched_rows.append([excel_row, source_id, region, country_code, country, recipe_name, ingredient_name, gram, "", "", "", "ingredient_name_not_in_match"])
            continue

        base_code = code_by_ingredient[lookup_key]
        used_code, status = resolve_idea_code(base_code, country_code, lcia_by_code)
        if used_code is None:
            unmatched_ingredient_count += 1
            agg["unmatched_rows"] += 1
            unmatched_rows.append([
                excel_row,
                source_id,
                region,
                country_code,
                country,
                recipe_name,
                ingredient_name,
                gram,
                base_code,
                replace_code_suffix(base_code, country_code),
                replace_code_suffix(base_code, "GLO"),
                status,
            ])
            continue

        impacts = [(gram / 1000.0) * factor for factor in lcia_by_code[used_code]]
        for i, value in enumerate(impacts):
            agg["total"][i] += value
        agg["matched_rows"] += 1
        matched_ingredient_count += 1

        matched_name = canonical_by_ingredient.get(lookup_key, lookup_key)
        if alias_used:
            matched_name = f"{matched_name} (alias)"
        ingredient_detail_rows.append([
            excel_row,
            source_id,
            region,
            country_code,
            country,
            recipe_name,
            ingredient_name,
            matched_name,
            gram,
            base_code,
            used_code,
            status,
            product_by_code.get(used_code, ""),
            *impacts,
        ])
    wb_ing.close()

    print("Reading nutrition master table and joining dish impacts...")
    wb_nut = load_workbook(nutrition_path, data_only=True, read_only=True)
    ws_nut = wb_nut["Nutrition"]
    nut_header = next(ws_nut.iter_rows(min_row=1, max_row=1, values_only=True))
    nut_headers = header_map(nut_header)
    nutrition_dish_count = 0
    joined_dish_count = 0
    missing_dish_count = 0
    id_joined_count = 0
    name_joined_count = 0
    seen_nutrition_keys: set[str] = set()
    duplicate_nutrition_keys: set[str] = set()

    main_rows: list[list[Any]] = [[
        *nut_header,
        "environment_join_status",
        "ingredient_rows",
        "matched_ingredient_rows",
        "unmatched_ingredient_rows",
        "source_idrecipe_list",
        *[f"total_{i + 1}_{impact_names[i]} [{impact_units[i]}]" for i in range(impact_count)],
        *[f"per_serving_{i + 1}_{impact_names[i]} [{impact_units[i]}]" for i in range(impact_count)],
    ]]
    summary_rows: list[list[Any]] = [[
        "idrecipe",
        "recipe_name",
        "country",
        "serving",
        "join_status",
        "ingredient_rows",
        "matched_ingredient_rows",
        "unmatched_ingredient_rows",
        "source_idrecipe_list",
        "total_climate_change_kg_CO2eq",
        "per_serving_climate_change_kg_CO2eq",
    ]]
    missing_rows: list[list[Any]] = [["idrecipe", "recipe_name", "country", "serving", "reason"]]

    for row in ws_nut.iter_rows(min_row=2, values_only=True):
        nutrition_dish_count += 1
        idrecipe = row[nut_headers["idrecipe"]]
        recipe_name = row[nut_headers["recipe name"]]
        country = row[nut_headers["country"]]
        serving = to_float(row[nut_headers["serving"]])
        name_key = recipe_key(country, recipe_name)
        source_id_key = normalize_text(idrecipe)

        if name_key in seen_nutrition_keys:
            duplicate_nutrition_keys.add(name_key)
        seen_nutrition_keys.add(name_key)

        join_key = ""
        join_mode = ""
        missing_reason = "recipe_not_found_in_Ingredient_Info"
        if source_id_key and source_id_key in dish_agg_by_source_id:
            join_key = dish_agg_by_source_id[source_id_key]
            join_mode = "idrecipe"
        elif name_key in dish_agg_by_name and name_key not in ambiguous_name_keys:
            join_key = dish_agg_by_name[name_key]
            join_mode = "country_recipe"
        elif name_key in ambiguous_name_keys:
            missing_reason = "ambiguous_country_recipe_without_idrecipe"

        total = vec(impact_count)
        ingredient_rows = matched_rows = unmatched_count = 0
        source_ids = ""
        status = missing_reason
        if join_key and join_key in dish_agg:
            joined_dish_count += 1
            if join_mode == "idrecipe":
                id_joined_count += 1
                status = "matched_by_idrecipe"
            else:
                name_joined_count += 1
                status = "matched_by_country_recipe"
            agg = dish_agg[join_key]
            total = list(agg["total"])
            ingredient_rows = int(agg["ingredient_rows"])
            matched_rows = int(agg["matched_rows"])
            unmatched_count = int(agg["unmatched_rows"])
            source_ids = ";".join(sorted(agg["source_ids"]))
            if matched_rows == 0:
                status = f"no_matched_ingredients_by_{join_mode}"
            elif unmatched_count > 0:
                status = f"partial_ingredient_match_by_{join_mode}"
        else:
            missing_dish_count += 1
            missing_rows.append([idrecipe, recipe_name, country, row[nut_headers["serving"]], status])

        per_serving = [value / serving if serving else None for value in total]
        main_rows.append([
            *row,
            status,
            ingredient_rows,
            matched_rows,
            unmatched_count,
            source_ids,
            *total,
            *per_serving,
        ])
        summary_rows.append([
            idrecipe,
            recipe_name,
            country,
            row[nut_headers["serving"]],
            status,
            ingredient_rows,
            matched_rows,
            unmatched_count,
            source_ids,
            total[0],
            per_serving[0],
        ])
    wb_nut.close()

    print("Writing output workbook...")
    out_wb = Workbook()
    del out_wb[out_wb.sheetnames[0]]
    readme_rows = [
        ["item", "value"],
        ["method", "Ingredient impacts = ingredient gram / 1000 * IDEA LCIA characterization factor. Dish totals are sums across ingredients. Per-serving values divide dish totals by Asia_recipe_nutrition serving."],
        ["code_lookup_priority", "recipe country code suffix -> GLO suffix -> base match code"],
        ["nutrition_file", str(nutrition_path)],
        ["ingredient_info_file", str(ingredient_info_path)],
        ["match_file", str(match_path)],
        ["LCIA_file", str(lcia_path)],
        ["nutrition_dishes", nutrition_dish_count],
        ["joined_dishes", joined_dish_count],
        ["idrecipe_joined_dishes", id_joined_count],
        ["country_recipe_fallback_joined_dishes", name_joined_count],
        ["missing_dishes", missing_dish_count],
        ["ingredient_records", ingredient_record_count],
        ["matched_ingredient_records", matched_ingredient_count],
        ["alias_resolved_ingredient_records", alias_resolved_count],
        ["unmatched_ingredient_records", unmatched_ingredient_count],
        ["invalid_or_blank_gram_records", invalid_gram_count],
        ["duplicate_match_names_skipped", duplicate_match_names],
        ["duplicate_nutrition_country_recipe_keys", len(duplicate_nutrition_keys)],
        ["ambiguous_ingredient_country_recipe_keys", len(ambiguous_name_keys)],
    ]
    for title, rows in [
        ("README", readme_rows),
        ("nutrition_plus_impact", main_rows),
        ("dish_GHG_summary", summary_rows),
        ("ingredient_detail", ingredient_detail_rows),
        ("unmatched_ingredients", unmatched_rows),
        ("missing_dishes", missing_rows),
    ]:
        ws = out_wb.create_sheet(title)
        append_rows(ws, rows)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    out_wb.save(output_path)
    print(f"Output: {output_path}")
    print(f"Dishes joined: {joined_dish_count} / {nutrition_dish_count}")
    print(f"Ingredient records matched: {matched_ingredient_count} / {ingredient_record_count}")
    print(f"Ingredient records unmatched: {unmatched_ingredient_count}")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Calculate dish-level IDEA LCIA impacts for Asia_recipe_nutrition.")
    parser.add_argument("--base-dir", default=".", help="Workspace base directory.")
    parser.add_argument("--nutrition-path", default="", help="Path to Asia_recipe_nutrition.xlsx.")
    parser.add_argument("--ingredient-info-path", default="", help="Path to Ingredient_Info.xlsx.")
    parser.add_argument("--match-path", default="", help="Path to match.xlsx.")
    parser.add_argument("--lcia-path", default="", help="Path to LCIA.xlsx.")
    parser.add_argument("--output-path", default="", help="Output workbook path.")
    return parser


if __name__ == "__main__":
    calculate(build_parser().parse_args())
